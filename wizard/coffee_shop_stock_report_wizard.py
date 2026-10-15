# -*- coding: utf-8 -*-
import base64
import io

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from odoo import models, fields
from datetime import datetime, timedelta


class CoffeeShopStockReportWizard(models.TransientModel):
    """
    Wizard to collect parameters for the Coffee Shop Stock Report Excel.
    """
    _name = 'coffee.shop.stock.report.wizard'
    _description = 'Coffee Shop Stock Report Wizard'

    filter_type = fields.Selection([
        ('all', 'All Products'),
        ('low_stock', 'Only Critical Stock (Below Threshold)')
    ], string='Stock Filter', required=True, default='all')

    def action_export_excel(self):
        """
        Generate the Coffee Shop Stock Report as an XLSX attachment.
        """
        self.ensure_one()

        report_data = self._get_stock_report_data()
        output = io.BytesIO()
        workbook = Workbook()
        header_fill = PatternFill(fill_type='solid', fgColor='6F4E37')
        header_font = Font(bold=True, color='FFFFFF')
        title_font = Font(bold=True, size=16)
        bold_font = Font(bold=True)
        danger_fill = PatternFill(fill_type='solid', fgColor='F4CCCC')

        summary_sheet = workbook.active
        summary_sheet.title = 'Summary'
        summary_sheet.column_dimensions['A'].width = 24
        summary_sheet.column_dimensions['B'].width = 26
        summary_sheet['A1'] = 'Laporan Stok Coffee Shop'
        summary_sheet['A1'].font = title_font
        summary_rows = [
            ('Perusahaan', report_data['company_name']),
            ('Filter', report_data['filter_type']),
            ('Total Produk', report_data['total_items']),
            ('Stok Kritis', report_data['low_stock_count']),
            ('Total Qty', report_data['total_qty']),
        ]
        for row, (label, value) in enumerate(summary_rows, start=2):
            summary_sheet.cell(row=row, column=1, value=label).font = bold_font
            value_cell = summary_sheet.cell(row=row, column=2, value=value)
            if label.startswith('Total'):
                value_cell.number_format = '#,##0.00'

        stock_sheet = workbook.create_sheet('Stock')
        headers = [
            'Product Name',
            'Product Category',
            'Qty',
            'Status',
        ]
        stock_sheet.append(headers)
        for column in range(1, len(headers) + 1):
            cell = stock_sheet.cell(row=1, column=column)
            cell.fill = header_fill
            cell.font = header_font
        for width_column, width in {
            'A': 30,
            'B': 24,
            'C': 16,
            'D': 16,
        }.items():
            stock_sheet.column_dimensions[width_column].width = width
        for product in report_data['products']:
            stock_sheet.append([
                product['name'],
                product['categ'],
                product['qty'],
                product['status'],
            ])
            if product['status'] == 'KRITIS':
                for cell in stock_sheet[stock_sheet.max_row]:
                    cell.fill = danger_fill
        for row in stock_sheet.iter_rows(min_row=2, min_col=3, max_col=3):
            for cell in row:
                cell.number_format = '#,##0.00'

        workbook.save(output)
        output.seek(0)

        current_date = (datetime.now() + timedelta(hours=7)).strftime("%d-%m-%Y %H:%M")
        filename = 'laporan_stok_%s_%s.xlsx' % (self.filter_type, current_date)
        attachment = self.env['ir.attachment'].create({
            'name': filename,
            'type': 'binary',
            'datas': base64.b64encode(output.read()),
            'res_model': self._name,
            'res_id': self.id,
            'mimetype': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        })
        return {
            'type': 'ir.actions.act_url',
            'url': '/web/content/%s?download=true' % attachment.id,
            'target': 'self',
        }

    def _get_stock_report_data(self):
        lang = self.env.context.get('lang') or 'en_US'
        query = """
            WITH quant_by_product AS (
                SELECT
                    sq.product_id,
                    SUM(
                        CASE
                            WHEN sl.usage = 'internal' THEN sq.quantity
                            ELSE 0
                        END
                    ) AS qty
                FROM stock_quant sq
                JOIN stock_location sl
                    ON sl.id = sq.location_id
                WHERE sq.company_id IS NULL
                OR sq.company_id = %s
                GROUP BY sq.product_id
            )
            SELECT
                COALESCE(pt.name->>%s, pt.name->>'en_US') AS product_name,
                COALESCE(pc.complete_name, pc.name, '-') AS product_category,
                COALESCE(qbp.qty, 0.0) AS qty,
                COALESCE(pt.low_stock_threshold, 0.0) AS minimum_stock,
                CASE
                    WHEN COALESCE(qbp.qty, 0.0) <= COALESCE(pt.low_stock_threshold, 0.0)
                        THEN 'KRITIS'
                    ELSE 'OK'
                END AS status
            FROM product_template pt
            JOIN product_product pp
                ON pp.product_tmpl_id = pt.id
            LEFT JOIN product_category pc
                ON pc.id = pt.categ_id
            LEFT JOIN quant_by_product qbp
                ON qbp.product_id = pp.id
            WHERE pt.active = TRUE
            AND pt.is_storable = TRUE
            AND (pt.company_id IS NULL OR pt.company_id = %s)
        """
        params = [self.env.company.id, lang, self.env.company.id]
        if self.filter_type == 'low_stock':
            query += """
              AND COALESCE(qbp.qty, 0.0) <= COALESCE(pt.low_stock_threshold, 0.0)
            """
        query += """
            ORDER BY COALESCE(pc.complete_name, pc.name, '-'), COALESCE(pt.name->>%s, pt.name->>'en_US')
        """
        params.append(lang)

        self.env.cr.execute(query, params)
        product_list = []
        total_qty = 0.0
        low_stock_count = 0
        for row in self.env.cr.dictfetchall():
            status = row['status']
            total_qty += row['qty'] or 0.0
            if status == 'KRITIS':
                low_stock_count += 1

            product_list.append({
                'name': row['product_name'] or '-',
                'categ': row['product_category'] or '-',
                'qty': row['qty'] or 0.0,
                'status': status,
            })

        return {
            'filter_type': 'Hanya Stok Menipis' if self.filter_type == 'low_stock' else 'Semua Produk',
            'products': product_list,
            'total_items': len(product_list),
            'total_qty': total_qty,
            'low_stock_count': low_stock_count,
            'company_name': self.env.company.name,
        }
