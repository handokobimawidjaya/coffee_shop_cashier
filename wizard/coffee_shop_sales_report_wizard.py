# -*- coding: utf-8 -*-
import base64
import io
from datetime import datetime, time

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from odoo import models, fields


class CoffeeShopSalesReportWizard(models.TransientModel):
    """
    Wizard to collect dates and category parameters for the Coffee Shop Sales Report Excel.
    """
    _name = 'coffee.shop.sales.report.wizard'
    _description = 'Coffee Shop Sales Report Wizard'

    date_start = fields.Date(
        string='Start Date',
        required=True,
        default=fields.Date.context_today
    )
    date_end = fields.Date(
        string='End Date',
        required=True,
        default=fields.Date.context_today
    )
    product_categ_ids = fields.Many2many(
        'product.category',
        string='Product Categories',
        help="Leave empty to include all categories."
    )

    def action_export_excel(self):
        """
        Generate the Coffee Shop Sales Report as an XLSX attachment.
        """
        self.ensure_one()

        report_data = self._get_sales_report_data()
        output = io.BytesIO()
        workbook = Workbook()
        header_fill = PatternFill(fill_type='solid', fgColor='6F4E37')
        header_font = Font(bold=True, color='FFFFFF')
        title_font = Font(bold=True, size=16)
        bold_font = Font(bold=True)

        summary_sheet = workbook.active
        summary_sheet.title = 'Summary'
        summary_sheet.column_dimensions['A'].width = 24
        summary_sheet.column_dimensions['B'].width = 28
        summary_sheet['A1'] = 'Laporan Penjualan Coffee Shop'
        summary_sheet['A1'].font = title_font
        summary_rows = [
            ('Perusahaan', report_data['company_name']),
            ('Tanggal Mulai', fields.Date.to_string(self.date_start)),
            ('Tanggal Selesai', fields.Date.to_string(self.date_end)),
            ('Kategori', report_data['categories']),
            ('Total Qty', report_data['total_qty']),
            ('Total Sales', report_data['total_sales']),
        ]
        for row, (label, value) in enumerate(summary_rows, start=2):
            summary_sheet.cell(row=row, column=1, value=label).font = bold_font
            value_cell = summary_sheet.cell(row=row, column=2, value=value)
            if label.startswith('Total'):
                value_cell.number_format = '#,##0.00'

        product_sheet = workbook.create_sheet('Sales')
        product_headers = [
            'Product Name',
            'Customer Name',
            'Product Category',
            'Payment Method',
            'Qty',
            'Total Sales',
        ]
        product_sheet.append(product_headers)
        for column in range(1, len(product_headers) + 1):
            cell = product_sheet.cell(row=1, column=column)
            cell.fill = header_fill
            cell.font = header_font
        for width_column, width in {
            'A': 30,
            'B': 24,
            'C': 24,
            'D': 20,
            'E': 16,
            'F': 18,
        }.items():
            product_sheet.column_dimensions[width_column].width = width
        for product in report_data['products']:
            product_sheet.append([
                product['name'],
                product['customer_name'],
                product['categ'],
                product['payment_method'],
                product['qty'],
                product['total_sales'],
            ])
        for row in product_sheet.iter_rows(min_row=2, min_col=5, max_col=6):
            for cell in row:
                cell.number_format = '#,##0.00'

        workbook.save(output)
        output.seek(0)

        filename = 'laporan_penjualan_%s_%s.xlsx' % (
            fields.Date.to_string(self.date_start),
            fields.Date.to_string(self.date_end),
        )
        attachment = self.env['ir.attachment'].create({
            'name': filename,
            'type': 'binary',
            'datas': base64.b64encode(output.read()),
            'res_model': self._name,
            'res_id': self.id,
            'mimetype': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        })
        return {
            'type': 'ir.actions.act_url',
            'url': '/web/content/%s?download=true' % attachment.id,
            'target': 'self',
        }

    def _get_sales_report_data(self):
        date_start = fields.Date.to_string(self.date_start)
        date_end = fields.Date.to_string(self.date_end)
        dt_start = datetime.combine(self.date_start, time.min)
        dt_end = datetime.combine(self.date_end, time.max)
        lang = self.env.context.get('lang') or 'en_US'

        query = """
            SELECT
                COALESCE(pt.name->>%s, pt.name->>'en_US') AS product_name,
                COALESCE(rp.complete_name , '') AS customer_name,
                COALESCE(pc.name, pc.complete_name, '-') AS product_category,
                COALESCE(payment_info.payment_methods, '-') AS payment_method,
                SUM(pol.qty) AS qty,
                SUM(pol.price_subtotal_incl) AS total_sales
            FROM pos_order_line pol
            JOIN pos_order po ON po.id = pol.order_id
            JOIN product_product pp ON pp.id = pol.product_id
            JOIN product_template pt ON pt.id = pp.product_tmpl_id
            LEFT JOIN product_category pc ON pc.id = pt.categ_id
            LEFT JOIN res_partner rp ON po.partner_id = rp.id
            LEFT JOIN (
                SELECT order_payments.pos_order_id,
                       string_agg(
                           order_payments.payment_method,
                           ', ' ORDER BY order_payments.payment_method
                       ) AS payment_methods
                  FROM (
                      SELECT DISTINCT pp.pos_order_id,
                             CASE
                                 WHEN ppm.is_cash_count THEN 'CASH'
                                 WHEN lower(COALESCE(ppm.name->>%s, ppm.name->>'en_US')) = 'qris'
                                     THEN 'QRIS'
                                 ELSE upper(COALESCE(ppm.name->>%s, ppm.name->>'en_US'))
                             END AS payment_method
                        FROM pos_payment pp
                        JOIN pos_payment_method ppm
                          ON ppm.id = pp.payment_method_id
                  ) order_payments
                 GROUP BY order_payments.pos_order_id
            ) payment_info ON payment_info.pos_order_id = po.id
            WHERE po.date_order >= %s
              AND po.date_order <= %s
              AND po.state IN ('paid', 'done', 'invoiced')
        """
        params = [
            lang,
            lang,
            lang,
            fields.Datetime.to_string(dt_start),
            fields.Datetime.to_string(dt_end),
        ]
        if self.product_categ_ids:
            query += """
              AND pt.categ_id = ANY(%s)
            """
            params.append(self.product_categ_ids.ids)

        query += """
            GROUP BY product_name, product_category, customer_name, payment_method
            ORDER BY product_name
        """
        self.env.cr.execute(query, params)

        products = []
        total_qty = 0.0
        total_sales = 0.0
        for row in self.env.cr.dictfetchall():
            qty = row['qty'] or 0.0
            sales = row['total_sales'] or 0.0
            total_qty += qty
            total_sales += sales
            products.append({
                'name': row['product_name'] or '-',
                'customer_name': row['customer_name'],
                'categ': row['product_category'] or '-',
                'payment_method': row['payment_method'] or '-',
                'qty': qty,
                'total_sales': sales,
            })

        return {
            'date_start': date_start,
            'date_end': date_end,
            'categories': ", ".join(self.product_categ_ids.mapped('name')) if self.product_categ_ids else "Semua Kategori",
            'products': products,
            'total_sales': total_sales,
            'total_qty': total_qty,
            'company_name': self.env.company.name,
        }
